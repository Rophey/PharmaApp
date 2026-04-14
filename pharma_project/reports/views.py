from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.http import HttpResponse
from .models import ReportTemplate, GeneratedReport, ProductionStats
from ebr.models import EBR
from quality.models import Deviation
import csv
import json
import datetime
import openpyxl
from openpyxl.styles import Font, Alignment


def report_list(request):
    if 'user_id' not in request.session or request.session.get('user_role') != 'директор':
        messages.error(request, 'Доступ только для директора')
        return redirect('users:director_dashboard')

    templates = ReportTemplate.objects.filter(is_active=True)
    recent_reports = GeneratedReport.objects.filter(
        generated_by=request.current_user
    ).order_by('-generated_at')[:20]

    context = {
        'user': request.current_user,
        'templates': templates,
        'recent_reports': recent_reports,
    }
    return render(request, 'reports/report_list.html', context)


def _generate_report_data(report_type, start_date, end_date):
    """Вспомогательная функция: генерирует данные отчёта (словарь)."""
    data = {}

    if report_type == 'production':
        ebrs = EBR.objects.filter(
            start_date__date__gte=start_date.date() if hasattr(start_date, 'date') else start_date,
            start_date__date__lte=end_date.date() if hasattr(end_date, 'date') else end_date,
        ).select_related('mbr__product', 'status', 'operator', 'signed_by')

        # Заголовки таблицы
        headers = ['Номер партии', 'Продукт', 'Статус', 'Дата начала', 'Оператор', 'Подписал', 'Дата решения']
        rows = []

        products_summary = {}
        total = 0
        completed = 0
        blocked = 0

        for ebr in ebrs:
            product_code = ebr.mbr.product.product_code
            product_name = ebr.mbr.product.product_name

            # Сводка по продуктам
            if product_code not in products_summary:
                products_summary[product_code] = {
                    'product_name': product_name,
                    'batches': 0, 'completed': 0, 'blocked': 0,
                }
            products_summary[product_code]['batches'] += 1

            # Строка таблицы
            operator_name = f"{ebr.operator.last_name} {ebr.operator.first_name}" if ebr.operator else '—'
            signed_name = f"{ebr.signed_by.last_name} {ebr.signed_by.first_name}" if ebr.signed_by else '—'
            decision_date = ebr.completion_date.strftime('%d.%m.%Y %H:%M') if ebr.completion_date else '—'

            rows.append([
                ebr.batch_number,
                f"{product_code} — {product_name}",
                ebr.status.status_name,
                ebr.start_date.strftime('%d.%m.%Y %H:%M'),
                operator_name,
                signed_name,
                decision_date,
            ])

            total += 1
            if ebr.status.status_name == 'Завершена':
                completed += 1
                products_summary[product_code]['completed'] += 1
            elif ebr.status.status_name == 'Заблокирована':
                blocked += 1
                products_summary[product_code]['blocked'] += 1

        data = {
            'summary_headers': ['Продукт', 'Партий', 'Завершено', 'Заблокировано', '% брака'],
            'summary_rows': [],
            'headers': headers,
            'rows': rows,
            'total': total,
            'completed': completed,
            'blocked': blocked,
            'defect_pct': round((blocked / total * 100), 1) if total > 0 else 0,
        }
        for code, pdata in products_summary.items():
            pct = round((pdata['blocked'] / pdata['batches'] * 100), 1) if pdata['batches'] > 0 else 0
            data['summary_rows'].append([
                f"{code} — {pdata['product_name']}",
                pdata['batches'],
                pdata['completed'],
                pdata['blocked'],
                f"{pct}%",
            ])

    elif report_type == 'deviation':
        deviations = Deviation.objects.filter(
            detected_at__date__gte=start_date.date() if hasattr(start_date, 'date') else start_date,
            detected_at__date__lte=end_date.date() if hasattr(end_date, 'date') else end_date,
        ).select_related('ebr', 'detected_by', 'resolved_by')

        headers = ['Партия', 'Параметр', 'Тип', 'Ожидаемое', 'Фактическое', 'Допуск', 'Обнаружил', 'Дата', 'Действие', 'Решил']
        rows = []
        types_count = {}

        for dev in deviations:
            types_count[dev.deviation_type] = types_count.get(dev.deviation_type, 0) + 1
            detected_by = f"{dev.detected_by.last_name} {dev.detected_by.first_name}" if dev.detected_by else '—'
            resolved_by = f"{dev.resolved_by.last_name} {dev.resolved_by.first_name}" if dev.resolved_by else '—'
            action = dev.get_action_taken_display() if dev.action_taken else '—'

            rows.append([
                dev.ebr.batch_number if dev.ebr else '—',
                dev.parameter_name,
                dev.get_deviation_type_display(),
                dev.expected_value,
                dev.actual_value,
                dev.tolerance,
                detected_by,
                dev.detected_at.strftime('%d.%m.%Y %H:%M'),
                action,
                resolved_by,
            ])

        data = {
            'total_deviations': deviations.count(),
            'by_type': types_count,
            'headers': headers,
            'rows': rows,
        }

    elif report_type == 'efficiency':
        ebrs = EBR.objects.filter(
            start_date__date__gte=start_date.date() if hasattr(start_date, 'date') else start_date,
            start_date__date__lte=end_date.date() if hasattr(end_date, 'date') else end_date,
        ).select_related('status')
        total = ebrs.count()
        completed = ebrs.filter(status__status_name='Завершена').count()
        blocked = ebrs.filter(status__status_name='Заблокирована').count()
        in_progress = ebrs.filter(status__status_name='В работе').count()
        waiting = ebrs.filter(status__status_name='В ожидании').count()

        data = {
            'total_batches': total,
            'completed': completed,
            'blocked': blocked,
            'in_progress': in_progress,
            'waiting': waiting,
            'efficiency': round((completed / total * 100), 1) if total > 0 else 0,
            'defect_rate': round((blocked / total * 100), 1) if total > 0 else 0,
        }

    return data


def generate_report(request, report_type):
    if 'user_id' not in request.session or request.session.get('user_role') != 'директор':
        return redirect('login')

    if request.method == 'POST':
        start_date = request.POST.get('start_date')
        end_date = request.POST.get('end_date')
        format = request.POST.get('format', 'excel')

        # Преобразуем строки в даты
        try:
            start = datetime.datetime.strptime(start_date, '%Y-%m-%d').date()
            end = datetime.datetime.strptime(end_date, '%Y-%m-%d').date()
        except (ValueError, TypeError):
            messages.error(request, 'Некорректный формат даты')
            return redirect('reports:report_list')

        try:
            template = ReportTemplate.objects.get(report_type=report_type)
        except ReportTemplate.DoesNotExist:
            template = ReportTemplate.objects.create(
                report_type=report_type,
                name=f'Отчёт {report_type}',
                is_active=True
            )

        # Генерируем данные в зависимости от типа отчёта
        data = {}

        if report_type == 'production':
            # Отчёт по выпуску продукции
            ebrs = EBR.objects.filter(
                start_date__date__gte=start,
                start_date__date__lte=end,
            ).select_related('mbr__product', 'status')

            products_data = {}
            total_units = 0
            total_defects = 0

            for ebr in ebrs:
                product_code = ebr.mbr.product.product_code
                if product_code not in products_data:
                    products_data[product_code] = {
                        'product_name': ebr.mbr.product.product_name,
                        'batches': 0,
                        'completed': 0,
                        'blocked': 0,
                        'defects': 0,
                    }

                products_data[product_code]['batches'] += 1
                if ebr.status.status_name == 'Завершена':
                    products_data[product_code]['completed'] += 1
                elif ebr.status.status_name == 'Заблокирована':
                    products_data[product_code]['blocked'] += 1
                    products_data[product_code]['defects'] += 1

                total_units += 1
                if ebr.status.status_name == 'Заблокирована':
                    total_defects += 1

            data = {
                'start_date': start_date,
                'end_date': end_date,
                'products': products_data,
                'total_units': total_units,
                'total_defects': total_defects,
                'defect_percentage': (total_defects / total_units * 100) if total_units > 0 else 0
            }

        elif report_type == 'deviation':
            # Отчёт по отклонениям
            deviations = Deviation.objects.filter(
                detected_at__date__gte=start,
                detected_at__date__lte=end
            )

            types_count = {}
            for dev in deviations:
                if dev.deviation_type not in types_count:
                    types_count[dev.deviation_type] = 0
                types_count[dev.deviation_type] += 1

            data = {
                'start_date': start_date,
                'end_date': end_date,
                'total_deviations': deviations.count(),
                'by_type': types_count,
                'deviations': deviations[:50]  # последние 50
            }

        # Сохраняем отчёт
        report = GeneratedReport.objects.create(
            template=template,
            generated_by=request.current_user,
            start_date=start,
            end_date=end,
            format=format,
            data=data
        )

        # Генерируем файл
        if format == 'excel':
            return generate_excel_report(report, data, report_type)
        elif format == 'csv':
            return generate_csv_report(report, data, report_type)

        messages.success(request, 'Отчёт сгенерирован')
        return redirect('reports:report_list')

    return redirect('reports:report_list')


def generate_excel_report(report, data, report_type):
    """Генерирует Excel-файл отчёта"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отчёт"

    # Заголовок
    ws['A1'] = f"Отчёт по {report_type} за период {data['start_date']} - {data['end_date']}"
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:E1')

    if report_type == 'production':
        # Заголовки таблицы
        headers = ['Продукт', 'Количество партий', 'Завершено', 'Заблокировано', 'Процент брака']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = Font(bold=True)

        # Данные
        row = 4
        for code, prod_data in data['products'].items():
            ws.cell(row=row, column=1, value=f"{code} - {prod_data['product_name']}")
            ws.cell(row=row, column=2, value=prod_data['batches'])
            ws.cell(row=row, column=3, value=prod_data['completed'])
            ws.cell(row=row, column=4, value=prod_data['blocked'])
            defect_pct = prod_data['defects'] / prod_data['batches'] * 100 if prod_data['batches'] > 0 else 0
            ws.cell(row=row, column=5, value=f"{defect_pct:.1f}%")
            row += 1

        # Итого
        ws.cell(row=row, column=1, value="ИТОГО").font = Font(bold=True)
        ws.cell(row=row, column=2, value=data['total_units']).font = Font(bold=True)
        ws.cell(row=row, column=5, value=f"{data['defect_percentage']:.1f}%").font = Font(bold=True)

    # Сохраняем
    import os
    from django.conf import settings
    filename = f"report_{report_type}_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(settings.MEDIA_ROOT, 'reports', filename)
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    wb.save(filepath)

    report.file = f"reports/{filename}"
    report.save()

    return redirect('reports:report_list')


def generate_csv_report(report, data, report_type):
    """Генерирует CSV-файл отчёта"""
    response = HttpResponse(content_type='text/csv')
    filename = f"report_{report_type}_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    writer = csv.writer(response)
    writer.writerow([f"Отчёт по {report_type}", data['start_date'], data['end_date']])
    writer.writerow([])

    if report_type == 'production':
        writer.writerow(['Продукт', 'Количество партий', 'Завершено', 'Заблокировано', 'Процент брака'])
        for code, prod_data in data['products'].items():
            defect_pct = prod_data['defects'] / prod_data['batches'] * 100 if prod_data['batches'] > 0 else 0
            writer.writerow([
                f"{code} - {prod_data['product_name']}",
                prod_data['batches'],
                prod_data['completed'],
                prod_data['blocked'],
                f"{defect_pct:.1f}%"
            ])

    elif report_type == 'deviation':
        writer.writerow(['Всего отклонений', data['total_deviations']])
        writer.writerow([])
        writer.writerow(['Тип', 'Количество'])
        for dtype, cnt in data.get('by_type', {}).items():
            writer.writerow([dtype, cnt])
        writer.writerow([])
        writer.writerow(['Параметр', 'Тип', 'Ожидаемое', 'Фактическое', 'Описание', 'Дата'])
        for dev_row in data.get('deviations', []):
            writer.writerow(dev_row)

    elif report_type == 'efficiency':
        writer.writerow(['Всего партий', data['total_batches']])
        writer.writerow(['Завершено', data['completed']])
        writer.writerow(['Заблокировано', data['blocked']])
        writer.writerow(['Эффективность', f"{data['efficiency']}%"])
        writer.writerow(['Брак', f"{data['defect_rate']}%"])

    return response


def download_report(request, report_id):
    if 'user_id' not in request.session:
        return redirect('login')

    report = get_object_or_404(GeneratedReport, id=report_id)

    if report.generated_by != request.current_user and request.session.get('user_role') != 'директор':
        messages.error(request, 'У вас нет доступа к этому отчёту')
        return redirect('reports:report_list')

    from django.http import FileResponse
    import os
    from django.conf import settings

    file_path = os.path.join(settings.MEDIA_ROOT, str(report.file))
    if os.path.exists(file_path):
        return FileResponse(open(file_path, 'rb'), as_attachment=True, filename=os.path.basename(file_path))

    messages.error(request, 'Файл не найден')
    return redirect('reports:report_list')